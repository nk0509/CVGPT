"""Step 1: read a CV file and turn it into structured data."""
import io
import json
import sys
from pathlib import Path

import pdfplumber
from docx import Document

from llm import ask_json

MAX_CV_CHARS = 12000  # keeps the prompt well inside the model's context window

PROMPT = """Extract information from this CV. Only use what the CV states; never guess.
Return JSON with exactly these keys:
{
  "name": string,
  "headline": string,
  "total_years_experience": number,
  "skills": [string],
  "tools": [string],
  "roles": [{"title": string, "company": string, "start": string, "end": string, "highlights": [string]}],
  "education": [{"degree": string, "institution": string, "year": string}],
  "certifications": [string]
}
Use "" or [] when something is missing.

CV:
\"\"\"
{cv}
\"\"\"
"""


def extract_text(source, filename=None):
    """Read plain text from a PDF, Word or text file.

    `source` can be a file path or an uploaded file object (bytes-like).
    """
    name = str(filename or source)
    suffix = Path(name).suffix.lower()
    if hasattr(source, "read"):
        source = io.BytesIO(source.read())

    if suffix == ".pdf":
        with pdfplumber.open(source) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n".join(pages).strip()
    if suffix == ".docx":
        doc = Document(source)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many CV templates put content in tables
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines).strip()
    if suffix in (".txt", ".md"):
        if isinstance(source, io.BytesIO):
            return source.getvalue().decode("utf-8", errors="ignore").strip()
        return Path(source).read_text(encoding="utf-8").strip()
    raise ValueError(f"Unsupported file type '{suffix}'. Use PDF, DOCX or TXT.")


def _as_list(value):
    if isinstance(value, list):
        return [v for v in value if v]
    return [value] if value else []


def parse_cv(text):
    """Ask the model to turn CV text into structured JSON, then tidy the result."""
    data = ask_json(PROMPT.replace("{cv}", text[:MAX_CV_CHARS]))
    try:
        years = float(data.get("total_years_experience") or 0)
    except (TypeError, ValueError):
        years = 0.0
    return {
        "name": str(data.get("name") or ""),
        "headline": str(data.get("headline") or ""),
        "total_years_experience": years,
        "skills": [str(s) for s in _as_list(data.get("skills"))],
        "tools": [str(s) for s in _as_list(data.get("tools"))],
        "roles": [r for r in _as_list(data.get("roles")) if isinstance(r, dict)],
        "education": [e for e in _as_list(data.get("education")) if isinstance(e, dict)],
        "certifications": [str(c) for c in _as_list(data.get("certifications"))],
    }


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python cv_parser.py cvs/your_cv.pdf")
        sys.exit(1)
    cv_path = Path(sys.argv[1])
    cv_text = extract_text(cv_path)
    print(f"Read {len(cv_text.split())} words from {cv_path.name}")
    result = parse_cv(cv_text)
    print(json.dumps(result, indent=2, ensure_ascii=False))
